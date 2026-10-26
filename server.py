import base64
import collections
import io
import json
import mimetypes
import os
import re
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlparse
from urllib.request import Request, urlopen

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import pypdf


def load_env_file() -> None:
    env_path = os.path.join(os.path.dirname(__file__), ".env")
    if not os.path.exists(env_path):
        return
    with open(env_path, encoding="utf-8") as env_file:
        for raw_line in env_file:
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


load_env_file()

HOST = os.getenv("HOST", "0.0.0.0")
PORT = int(os.getenv("PORT", "8080"))
OLLAMA_URL = os.getenv("OLLAMA_URL", "http://localhost:11434")
MODEL_NAME = os.getenv("MODEL_NAME", "llama3.2:3b")
PRESENTATION_MODEL = os.getenv("PRESENTATION_MODEL", "qwen2.5:7b")
DEMO_MODE = os.getenv("DEMO_MODE", "false").lower() == "true"
MAX_CONCURRENT_REQUESTS = int(os.getenv("MAX_CONCURRENT_REQUESTS", "10"))
OLLAMA_NUM_PARALLEL = int(os.getenv("OLLAMA_NUM_PARALLEL", "5"))

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "generated")
DIST_DIR = os.path.join(os.path.dirname(__file__), "dist")
os.makedirs(OUTPUT_DIR, exist_ok=True)


class MetricsTracker:
    def __init__(self, max_history: int = 50, max_buckets: int = 15):
        self.lock = threading.Lock()
        self.active_requests = 0
        self.peak_active_requests = 0
        self.total_requests = 0
        self.successful_requests = 0
        self.failed_requests = 0
        self.total_latency_ms = 0
        self.min_latency_ms: int | None = None
        self.max_latency_ms = 0
        self.total_presentations = 0
        self.total_documents = 0
        self.active_users: dict[str, float] = {}
        self.peak_active_users = 0
        self.use_case_counts: dict[str, int] = collections.defaultdict(int)
        self.recent_requests: collections.deque = collections.deque(maxlen=max_history)
        self.minute_buckets: collections.OrderedDict[str, dict] = collections.OrderedDict()
        self.max_buckets = max_buckets
        self.req_id_counter = 0

        # Pre-seed baseline timeline points
        now = time.time()
        for i in range(5, -1, -1):
            t_str = time.strftime("%H:%M", time.localtime(now - i * 60))
            self.minute_buckets[t_str] = {
                "time": t_str,
                "requests": 0,
                "total_latency": 0,
                "errors": 0,
                "peak_concurrent": 0,
                "active_users": 1,
            }

    def touch_user(self, client_id: str) -> None:
        if not client_id:
            return
        now = time.time()
        with self.lock:
            self.active_users[client_id] = now
            self._clean_users_locked(now)

    def _clean_users_locked(self, now: float) -> None:
        cutoff = now - 300  # 5 minutes
        self.active_users = {k: v for k, v in self.active_users.items() if v >= cutoff}
        count = len(self.active_users)
        if count > self.peak_active_users:
            self.peak_active_users = count

    def start_request(self, client_id: str) -> tuple[int, float]:
        now = time.time()
        with self.lock:
            self.active_requests += 1
            if self.active_requests > self.peak_active_requests:
                self.peak_active_requests = self.active_requests
            if client_id:
                self.active_users[client_id] = now
            self._clean_users_locked(now)
            self.req_id_counter += 1
            req_id = self.req_id_counter
        return req_id, now

    def end_request(
        self,
        req_id: int,
        client_id: str,
        use_case: str,
        latency_ms: int,
        status_code: int,
        mode: str,
        is_presentation: bool,
    ) -> None:
        now = time.time()
        minute_key = time.strftime("%H:%M", time.localtime(now))
        is_success = 200 <= status_code < 400

        with self.lock:
            self.active_requests = max(0, self.active_requests - 1)
            self.total_requests += 1
            if is_success:
                self.successful_requests += 1
            else:
                self.failed_requests += 1

            self.total_latency_ms += latency_ms
            if self.min_latency_ms is None or latency_ms < self.min_latency_ms:
                self.min_latency_ms = latency_ms
            if latency_ms > self.max_latency_ms:
                self.max_latency_ms = latency_ms

            if is_presentation:
                self.total_presentations += 1
                self.use_case_counts["Presentaciones"] += 1
            elif use_case:
                self.use_case_counts[use_case] += 1

            if client_id:
                self.active_users[client_id] = now
            self._clean_users_locked(now)

            self.recent_requests.appendleft(
                {
                    "id": f"req-{req_id}",
                    "timestamp": time.strftime("%H:%M:%S", time.localtime(now)),
                    "client_id": client_id[-6:] if client_id else "anon",
                    "use_case": "Presentación" if is_presentation else use_case,
                    "latency_ms": latency_ms,
                    "status": status_code,
                    "mode": mode,
                    "is_presentation": is_presentation,
                }
            )

            bucket = self.minute_buckets.get(minute_key)
            if not bucket:
                if len(self.minute_buckets) >= self.max_buckets:
                    self.minute_buckets.popitem(last=False)
                bucket = {
                    "time": minute_key,
                    "requests": 0,
                    "total_latency": 0,
                    "errors": 0,
                    "peak_concurrent": self.active_requests + 1,
                    "active_users": max(1, len(self.active_users)),
                }
                self.minute_buckets[minute_key] = bucket

            bucket["requests"] += 1
            bucket["total_latency"] += latency_ms
            if not is_success:
                bucket["errors"] += 1
            bucket["peak_concurrent"] = max(bucket["peak_concurrent"], self.active_requests + 1)
            bucket["active_users"] = max(bucket["active_users"], len(self.active_users))

    def get_summary(self) -> dict:
        now = time.time()
        with self.lock:
            self._clean_users_locked(now)
            avg_latency = (
                round(self.total_latency_ms / self.total_requests)
                if self.total_requests > 0
                else 0
            )
            success_rate = (
                round((self.successful_requests / self.total_requests) * 100, 1)
                if self.total_requests > 0
                else 100.0
            )

            timeline = []
            for minute_key, bucket in self.minute_buckets.items():
                reqs = bucket["requests"]
                avg_l = round(bucket["total_latency"] / reqs) if reqs > 0 else 0
                timeline.append(
                    {
                        "time": bucket["time"],
                        "requests": reqs,
                        "avg_latency_ms": avg_l,
                        "peak_concurrent": bucket["peak_concurrent"],
                        "active_users": bucket["active_users"],
                        "errors": bucket["errors"],
                    }
                )

            active_users_count = max(len(self.active_users), 1 if self.active_requests > 0 else 0)
            peak_users_count = max(self.peak_active_users, active_users_count)

            return {
                "summary": {
                    "active_requests": self.active_requests,
                    "peak_active_requests": self.peak_active_requests,
                    "active_users_5m": active_users_count,
                    "peak_active_users": peak_users_count,
                    "total_requests": self.total_requests,
                    "successful_requests": self.successful_requests,
                    "failed_requests": self.failed_requests,
                    "success_rate_percent": success_rate,
                    "avg_latency_ms": avg_latency,
                    "min_latency_ms": self.min_latency_ms if self.min_latency_ms is not None else 0,
                    "max_latency_ms": self.max_latency_ms,
                    "total_presentations": self.total_presentations,
                    "total_documents": self.total_documents,
                    "max_concurrency_slots": MAX_CONCURRENT_REQUESTS,
                    "ollama_parallel_slots": OLLAMA_NUM_PARALLEL,
                    "mode": "demo" if DEMO_MODE else "ollama",
                    "model": MODEL_NAME,
                    "presentation_model": PRESENTATION_MODEL,
                },
                "by_use_case": dict(self.use_case_counts),
                "timeline": timeline,
                "recent_requests": list(self.recent_requests),
            }


metrics_tracker = MetricsTracker()
concurrency_semaphore = threading.BoundedSemaphore(MAX_CONCURRENT_REQUESTS)


def ollama_chat(
    message: str,
    use_case: str,
    model: str = MODEL_NAME,
    document_context: dict | None = None,
    history: list[dict] | None = None,
) -> str:
    system_prompt = (
        "Eres CIRCE Local AI, el asistente inteligente de élite de Fundación CIRCE (Centro Tecnológico).\n"
        f"Área actual de trabajo: {use_case}.\n"
        "CAPACIDADES CLAVE:\n"
        "- Capacidad analítica rigurosa en transición energética, sostenibilidad, industria e innovación tecnológica.\n"
        "- Cuando redactes documentos o informes técnicos, usa formato Markdown enriquecido con títulos (#), "
        "subtítulos (##, ###), listas estructuradas con viñetas, tablas comparativas y un apartado claro de conclusiones ejecutivas y pasos de acción.\n"
        "- Si el usuario solicita un documento o exportación a Word, redacta el contenido con estructura formal ejecutiva.\n"
        "- Si se adjunta un documento, analiza detalladamente sus datos, citas numéricas y metodología sin inventar hechos.\n"
        "- Si el usuario pide una presentación o diapositivas, resume en ideas clave con métricas cuantitativas e impacto directo.\n"
        "- Responde siempre en español, con lenguaje profesional, preciso y directo."
    )

    ollama_messages = [{"role": "system", "content": system_prompt}]

    # Include recent conversation history for multi-turn awareness
    if history:
        for turn in history[-6:]:
            role = turn.get("role", "user")
            content = turn.get("content", "")
            if role in ("user", "assistant") and content:
                ollama_messages.append({"role": role, "content": content})

    # Prepare current user content
    user_content = message
    if document_context and document_context.get("text"):
        doc_name = document_context.get("filename", "documento")
        doc_words = document_context.get("word_count", 0)
        doc_pages = document_context.get("page_count", 1)
        user_content = (
            f"[CONTEXTO DOCUMENTAL ADJUNTO: '{doc_name}' ({doc_pages} páginas, {doc_words} palabras)]\n"
            "--- INICIO CONTENIDO DOCUMENTO ---\n"
            f"{document_context['text']}\n"
            "--- FIN CONTENIDO DOCUMENTO ---\n\n"
            f"SOLICITUD DEL USUARIO:\n{message}"
        )

    # Avoid duplicate message if already at the end of history
    if not (ollama_messages and ollama_messages[-1].get("role") == "user" and ollama_messages[-1].get("content") == user_content):
        ollama_messages.append({"role": "user", "content": user_content})

    payload = json.dumps(
        {
            "model": model,
            "stream": False,
            "messages": ollama_messages,
        }
    ).encode("utf-8")
    request = Request(f"{OLLAMA_URL}/api/chat", data=payload, headers={"Content-Type": "application/json"})
    with urlopen(request, timeout=35) as response:
        return json.loads(response.read().decode("utf-8"))["message"]["content"]


def is_presentation_intent(message: str) -> bool:
    msg = message.lower()
    terms = (
        "presentación", "presentacion", "powerpoint", "diapositivas", "slides", "pptx",
        "conviértelo en presentación", "conviertelo en presentacion", "conviértelo a presentación",
        "conviertelo a presentacion", "pásalo a presentación", "pasalo a presentacion",
        "haz una presentación", "haz una presentacion", "hazme una presentación", "hazme una presentacion",
        "crea una presentación", "crea una presentacion", "haz diapositivas", "crea diapositivas",
        "generar presentación", "generar presentacion"
    )
    return any(term in msg for term in terms)


def is_export_word_intent(message: str) -> bool:
    msg = message.lower()
    terms = (
        "exporta a word", "expórtalo a word", "exportalo a word", "exportar a word",
        "pasar a word", "pásalo a word", "pasalo a word", "guarda en word", "guardar en word",
        "en formato word", "a formato word", "descargar en word", "descargar docx", "en docx",
        "informe word", "documento word", "redacta un informe", "crea un informe", "generar informe",
        "redáctalo en word", "redactalo en word", "crear informe word"
    )
    return any(term in msg for term in terms)


# ==========================================
# GESTIÓN Y EXTRACCIÓN DE DOCUMENTOS
# ==========================================

def extract_document_text(file_bytes: bytes, filename: str, mime_type: str = "") -> dict:
    ext = os.path.splitext(filename.lower())[1]
    text = ""
    page_count = 1

    if ext == ".pdf" or "pdf" in mime_type.lower():
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            page_count = len(reader.pages)
            pages_text = []
            for i, page in enumerate(reader.pages):
                extracted = page.extract_text() or ""
                if extracted.strip():
                    pages_text.append(f"--- Página {i+1} ---\n{extracted.strip()}")
            text = "\n\n".join(pages_text)
        except Exception as e:
            text = f"[Error al extraer texto del archivo PDF: {str(e)}]"

    elif ext in (".docx", ".doc") or "word" in mime_type.lower():
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_cells:
                        parts.append(" | ".join(row_cells))
            text = "\n\n".join(parts)
            page_count = max(1, round(len(parts) / 8))
        except Exception as e:
            text = f"[Error al extraer texto del documento Word: {str(e)}]"

    else:
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode("latin-1")
            except Exception as e:
                text = f"[Error al leer archivo de texto: {str(e)}]"
        page_count = max(1, round(len(text) / 2500))

    clean_text = text.strip()
    words = len(clean_text.split())
    chars = len(clean_text)
    preview = clean_text[:450] + ("..." if chars > 450 else "")

    safe_text = clean_text[:35000]
    if len(clean_text) > 35000:
        safe_text += f"\n\n[Texto truncado a 35.000 caracteres para el modelo. Total original: {chars} caracteres, {words} palabras]."

    return {
        "filename": filename,
        "extension": ext,
        "page_count": page_count,
        "word_count": words,
        "char_count": chars,
        "preview": preview,
        "text": safe_text,
    }


def _add_formatted_runs(paragraph, text: str, default_font="Calibri", default_size=11):
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.font.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.italic = True
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = DocxPt(9.5)
            run.font.color.rgb = DocxRGBColor(0x10, 0x85, 0x54)
        else:
            run = paragraph.add_run(part)

        if not (part.startswith("`") and part.endswith("`")):
            run.font.name = default_font
            run.font.size = DocxPt(default_size)


def create_docx_report(title: str, markdown_content: str, topic: str = "", metadata: dict | None = None) -> str:
    metadata = metadata or {}
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = DocxInches(1.0)
        section.bottom_margin = DocxInches(1.0)
        section.left_margin = DocxInches(1.0)
        section.right_margin = DocxInches(1.0)

        header = section.header
        hp = header.paragraphs[0]
        hp.text = "FUNDACIÓN CIRCE  ·  CENTRO TECNOLÓGICO  |  LOCAL AI LAB"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if hp.runs:
            hp.runs[0].font.size = DocxPt(8.5)
            hp.runs[0].font.color.rgb = DocxRGBColor(0x77, 0x82, 0x7C)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Fundación CIRCE — Confidencialidad y Privacidad Garantizadas — Procesado 100% Local"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if fp.runs:
            fp.runs[0].font.size = DocxPt(8)
            fp.runs[0].font.color.rgb = DocxRGBColor(0x99, 0xA2, 0x9C)

    doc_title = title if title and len(title) > 3 else (topic if topic else "Informe Técnico CIRCE")
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = DocxPt(0)
    title_p.paragraph_format.space_after = DocxPt(4)
    run_title = title_p.add_run(doc_title)
    run_title.font.name = "Calibri"
    run_title.font.size = DocxPt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = DocxRGBColor(0x10, 0x85, 0x54)

    date_str = time.strftime("%d de %B de %Y", time.localtime())
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = DocxPt(14)
    run_sub = sub_p.add_run(f"Informe técnico estructurado · {date_str}")
    run_sub.font.name = "Calibri"
    run_sub.font.size = DocxPt(10.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = DocxRGBColor(0x77, 0x82, 0x7C)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = DocxInches(6.5)
    cell = table.cell(0, 0)

    try:
        xml_shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F6F2"/>')
        cell._tc.get_or_add_tcPr().append(xml_shd)
    except Exception:
        pass

    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = DocxPt(6)
    cp.paragraph_format.space_after = DocxPt(6)

    r1 = cp.add_run("Ámbito: ")
    r1.font.bold = True
    r1.font.size = DocxPt(9.5)
    cp.add_run(f"{topic or 'Investigación Aplicada'}   |   ")

    r2 = cp.add_run("Seguridad: ")
    r2.font.bold = True
    r2.font.size = DocxPt(9.5)
    cp.add_run("100% Local (Garantía de privacidad CIRCE)\n")

    if metadata.get("doc_name"):
        r3 = cp.add_run("Documento Fuente: ")
        r3.font.bold = True
        r3.font.size = DocxPt(9.5)
        cp.add_run(f"{metadata['doc_name']} ({metadata.get('doc_words', 0)} palabras)")

    doc.add_paragraph().paragraph_format.space_after = DocxPt(10)

    lines = markdown_content.splitlines()
    in_code_block = False
    code_lines = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                code_text = "\n".join(code_lines)
                code_p = doc.add_paragraph()
                code_p.paragraph_format.left_indent = DocxInches(0.4)
                code_p.paragraph_format.space_after = DocxPt(6)
                code_run = code_p.add_run(code_text)
                code_run.font.name = "Consolas"
                code_run.font.size = DocxPt(9.5)
                code_run.font.color.rgb = DocxRGBColor(0x23, 0x69, 0xA5)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if not stripped:
            continue

        if stripped.startswith("### "):
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = DocxPt(10)
            hp.paragraph_format.space_after = DocxPt(3)
            r = hp.add_run(stripped[4:])
            r.font.name = "Calibri"
            r.font.size = DocxPt(12)
            r.font.bold = True
            r.font.color.rgb = DocxRGBColor(0x17, 0x35, 0x2B)
        elif stripped.startswith("## "):
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = DocxPt(14)
            hp.paragraph_format.space_after = DocxPt(4)
            r = hp.add_run(stripped[3:])
            r.font.name = "Calibri"
            r.font.size = DocxPt(14)
            r.font.bold = True
            r.font.color.rgb = DocxRGBColor(0x23, 0x69, 0xA5)
        elif stripped.startswith("# "):
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = DocxPt(16)
            hp.paragraph_format.space_after = DocxPt(5)
            r = hp.add_run(stripped[2:])
            r.font.name = "Calibri"
            r.font.size = DocxPt(16)
            r.font.bold = True
            r.font.color.rgb = DocxRGBColor(0x10, 0x85, 0x54)
        elif stripped.startswith(("- ", "* ", "• ")):
            bullet_text = stripped[2:].strip()
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.space_before = DocxPt(2)
            bp.paragraph_format.space_after = DocxPt(3)
            _add_formatted_runs(bp, bullet_text)
        elif re.match(r"^\d+\.\s+", stripped):
            num_match = re.match(r"^(\d+\.\s+)(.*)$", stripped)
            np = doc.add_paragraph(style="List Number")
            np.paragraph_format.space_before = DocxPt(2)
            np.paragraph_format.space_after = DocxPt(3)
            _add_formatted_runs(np, num_match.group(2) if num_match else stripped)
        elif stripped.startswith("> "):
            qp = doc.add_paragraph()
            qp.paragraph_format.left_indent = DocxInches(0.4)
            qp.paragraph_format.space_before = DocxPt(4)
            qp.paragraph_format.space_after = DocxPt(6)
            r = qp.add_run(stripped[2:])
            r.font.italic = True
            r.font.color.rgb = DocxRGBColor(0x46, 0x62, 0x54)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = DocxPt(2)
            p.paragraph_format.space_after = DocxPt(6)
            p.paragraph_format.line_spacing = 1.15
            _add_formatted_runs(p, stripped)

    safe_title = re.sub(r"[^a-zA-Z0-9áéíóúÁÉÍÓÚñÑ_-]+", "-", doc_title).strip("-").lower()[:45] or "informe-circe"
    filename = f"{safe_title}-{int(time.time())}.docx"
    doc.save(os.path.join(OUTPUT_DIR, filename))
    return filename


# ==========================================
# GENERACIÓN AVANZADA DE PRESENTACIONES PPTX (ESTILO CONSULTORÍA EJECUTIVA)
# ==========================================

THEMES = {
    "circe": {
        "background": "F3F7F4",
        "ink": "0F2E23",         # Deep Forest Green
        "accent": "108554",      # Vibrant Emerald Green
        "secondary": "D2E8DB",   # Pale Sage Mint
        "highlight": "D97706",   # Gold Amber
        "muted": "4E6458",       # Slate Olive
        "card": "FFFFFF",
    },
    "azul": {
        "background": "F0F5FA",
        "ink": "0F284E",         # Midnight Navy
        "accent": "1E70BF",      # Cobalt Blue
        "secondary": "D0E1F4",   # Ice Blue
        "highlight": "E67E22",   # Sunset Orange
        "muted": "4D627B",       # Slate Navy
        "card": "FFFFFF",
    },
    "calida": {
        "background": "FDF8F3",
        "ink": "341C14",         # Dark Espresso
        "accent": "C0392B",      # Terracotta Red
        "secondary": "F6DFD5",   # Peach Mist
        "highlight": "D97706",   # Warm Amber
        "muted": "6B5047",       # Muted Coffee
        "card": "FFFFFF",
    },
    "purpura": {
        "background": "F9F5FB",
        "ink": "2B1343",         # Obsidian Violet
        "accent": "8E44AD",      # Royal Purple
        "secondary": "EBD7F7",   # Lilac Mist
        "highlight": "E67E22",   # Vivid Orange
        "muted": "5E4870",       # Slate Plum
        "card": "FFFFFF",
    },
    "monocroma": {
        "background": "F4F5F6",
        "ink": "1A1F24",         # Charcoal Black
        "accent": "3E4B59",      # Steel Slate
        "secondary": "DDE2E6",   # Silver Mist
        "highlight": "52606D",   # Titanium
        "muted": "616E7C",       # Muted Slate
        "card": "FFFFFF",
    },
}


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_text(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    size: int,
    fill: str,
    bold: bool = False,
    font: str = "Aptos",
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color(fill)
    return box


def create_presentation(topic: str, slides: list[dict], preferences: dict) -> str:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    theme = THEMES.get(preferences.get("palette", "circe"), THEMES["circe"])
    density = preferences.get("density", "medio")
    bullet_size = {"poco": 15, "medio": 13, "mucho": 11}.get(density, 13)

    total_slides = len(slides)

    for index, slide_data in enumerate(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = color(theme["background"])

        layout = str(slide_data.get("layout", "")).lower()

        # ============================================================
        # 1. PORTADA EJECUTIVA DE GRAN IMPACTO (COVER SLIDE)
        # ============================================================
        if index == 0 or layout == "cover":
            # Left Dark Hero Panel (60% width)
            left_hero = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(7.8), Inches(7.5))
            left_hero.fill.solid()
            left_hero.fill.fore_color.rgb = color(theme["ink"])
            left_hero.line.fill.background()

            # Logo Badge Box
            logo_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.8), Inches(0.55), Inches(0.55))
            logo_box.fill.solid()
            logo_box.fill.fore_color.rgb = color(theme["accent"])
            logo_box.line.fill.background()
            add_text(slide, "C", 0.9, 0.82, 0.55, 0.5, 20, "FFFFFF", True, align=PP_ALIGN.CENTER)

            add_text(slide, "FUNDACIÓN CIRCE", 1.6, 0.82, 5.5, 0.3, 14, "FFFFFF", True)
            add_text(slide, "CENTRO TECNOLÓGICO · LOCAL AI LAB", 1.6, 1.15, 5.5, 0.25, 9, theme["secondary"], True)

            # Category Pill Tag
            cat_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(2.05), Inches(3.4), Inches(0.35))
            cat_pill.fill.solid()
            cat_pill.fill.fore_color.rgb = color(theme["accent"])
            cat_pill.line.fill.background()
            add_text(slide, "INFORME ESTRATÉGICO Y TECNOLÓGICO", 0.9, 2.1, 3.4, 0.28, 9, "FFFFFF", True, align=PP_ALIGN.CENTER)

            # Huge Main Title
            title_text = str(slide_data.get("title", topic))
            add_text(slide, title_text, 0.9, 2.65, 6.2, 2.2, 34, "FFFFFF", True, font="Georgia")

            # Subtitle
            sub_text = str(slide_data.get("subtitle", "Síntesis ejecutiva y análisis cuantitativo"))
            add_text(slide, sub_text, 0.9, 4.95, 6.2, 1.1, 15, theme["secondary"])

            # Confidentiality Ribbon
            add_text(slide, "PROCESAMIENTO 100% LOCAL  ·  CONFIDENCIALIDAD CIRCE GARANTIZADA", 0.9, 6.45, 6.2, 0.35, 9, theme["highlight"], True)

            # Right Side Executive Overview Card
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.4), Inches(1.2), Inches(4.1), Inches(5.1))
            card.fill.solid()
            card.fill.fore_color.rgb = color(theme["card"])
            card.line.color.rgb = color(theme["secondary"])

            top_tab = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.4), Inches(1.2), Inches(4.1), Inches(0.15))
            top_tab.fill.solid()
            top_tab.fill.fore_color.rgb = color(theme["accent"])
            top_tab.line.fill.background()

            add_text(slide, "ESTRUCTURA DEL ANÁLISIS", 8.8, 1.6, 3.3, 0.3, 10, theme["highlight"], True)
            add_text(slide, "Contenido Ejecutivo", 8.8, 1.95, 3.3, 0.4, 15, theme["ink"], True)

            highlights = [
                ("01", "Diagnóstico Técnico & Retos"),
                ("02", "Métricas Cuantitativas & Impacto"),
                ("03", "Pilares Estratégicos CIRCE"),
                ("04", "Plan de Fases & Hoja de Ruta"),
            ]
            for hi, (num, htext) in enumerate(highlights):
                top_h = 2.6 + hi * 0.72
                dot = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(top_h), Inches(0.45), Inches(0.35))
                dot.fill.solid()
                dot.fill.fore_color.rgb = color(theme["secondary"])
                dot.line.fill.background()
                add_text(slide, num, 8.8, top_h + 0.05, 0.45, 0.28, 9, theme["ink"], True, align=PP_ALIGN.CENTER)
                add_text(slide, htext, 9.4, top_h + 0.05, 2.8, 0.4, 11.5, theme["ink"], True)

            date_str = time.strftime("%d de %B de %Y", time.localtime())
            add_text(slide, f"Fundación CIRCE · {date_str}", 8.8, 5.75, 3.3, 0.3, 9, theme["muted"])
            continue

        # ============================================================
        # 2. CIERRE Y PRÓXIMOS PASOS (CLOSING SLIDE)
        # ============================================================
        if slide_data.get("kind") == "closing" or layout == "closing" or index == total_slides - 1:
            left_closing = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(5.0), Inches(7.5))
            left_closing.fill.solid()
            left_closing.fill.fore_color.rgb = color(theme["ink"])
            left_closing.line.fill.background()

            add_text(slide, "FUNDACIÓN CIRCE", 0.9, 0.9, 3.5, 0.35, 14, "FFFFFF", True)
            add_text(slide, "CENTRO TECNOLÓGICO", 0.9, 1.25, 3.5, 0.25, 9, theme["secondary"], True)

            add_text(slide, "Conocimiento\nque se convierte\nen acción.", 0.9, 2.5, 3.5, 2.2, 28, "FFFFFF", True, font="Georgia")
            add_text(slide, "www.fcirce.es  ·  Zaragoza, España\nLocal AI Lab · Innovación Aplicada", 0.9, 5.8, 3.5, 0.8, 10, theme["secondary"])

            # Right Side Action Plan
            add_text(slide, "PLAN DE ACCIÓN Y PRÓXIMOS PASOS", 5.6, 0.9, 6.8, 0.35, 11, theme["highlight"], True)
            add_text(slide, str(slide_data.get("title", "Conclusiones y Hoja de Ruta")), 5.6, 1.3, 6.8, 0.7, 26, theme["ink"], True)
            add_text(slide, str(slide_data.get("subtitle", "Actuaciones inmediatas para la toma de decisiones")), 5.6, 2.05, 6.8, 0.4, 13, theme["muted"])

            bullets = slide_data.get("bullets", [
                "Validar los requerimientos técnicos y alcance con los equipos especializados.",
                "Priorizar los casos de uso con mayor retorno energético y operativo.",
                "Establecer el calendario de hitos y el comité de seguimiento.",
            ])
            for bi, b in enumerate(bullets[:3]):
                top_b = 2.85 + bi * 1.25
                b_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(top_b), Inches(6.8), Inches(1.05))
                b_card.fill.solid()
                b_card.fill.fore_color.rgb = color(theme["card"])
                b_card.line.color.rgb = color(theme["secondary"])

                num_badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.85), Inches(top_b + 0.25), Inches(0.55), Inches(0.55))
                num_badge.fill.solid()
                num_badge.fill.fore_color.rgb = color(theme["accent"])
                num_badge.line.fill.background()
                add_text(slide, f"✓", 5.85, top_b + 0.28, 0.55, 0.45, 14, "FFFFFF", True, align=PP_ALIGN.CENTER)

                add_text(slide, str(b), 6.6, top_b + 0.2, 5.6, 0.7, 13, theme["ink"], True)

            add_text(slide, "CIRCE  ·  Documento Confidencial para Uso Interno", 5.6, 6.75, 6.8, 0.3, 9, theme["muted"])
            continue

        # ============================================================
        # MASTER HEADER & FOOTER FOR ALL CONTENT SLIDES
        # ============================================================
        top_hairline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
        top_hairline.fill.solid()
        top_hairline.fill.fore_color.rgb = color(theme["accent"])
        top_hairline.line.fill.background()

        # Category Tag Pill
        cat_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.42), Inches(2.9), Inches(0.28))
        cat_box.fill.solid()
        cat_box.fill.fore_color.rgb = color(theme["secondary"])
        cat_box.line.fill.background()
        cat_text = str(slide_data.get("category", "CIRCE · INNOVACIÓN TECNOLÓGICA")).upper()
        add_text(slide, cat_text, 0.9, 0.45, 2.9, 0.22, 8.5, theme["accent"], True, align=PP_ALIGN.CENTER)

        # Slide Number Badge Top Right
        num_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.2), Inches(0.42), Inches(1.2), Inches(0.28))
        num_pill.fill.solid()
        num_pill.fill.fore_color.rgb = color(theme["card"])
        num_pill.line.color.rgb = color(theme["secondary"])
        add_text(slide, f"{index:02d} / {total_slides:02d}", 11.2, 0.45, 1.2, 0.22, 8.5, theme["accent"], True, align=PP_ALIGN.CENTER)

        # Slide Main Title & Subtitle Lead-in
        add_text(slide, str(slide_data.get("title", "Aspecto Clave")), 0.9, 0.8, 9.8, 0.65, 24, theme["ink"], True, font="Georgia")
        sub_lead = str(slide_data.get("subtitle", "Principales evidencias, datos y consideraciones estratégicas."))
        add_text(slide, sub_lead, 0.9, 1.45, 10.5, 0.35, 11.5, theme["muted"])

        bullets = slide_data.get("bullets", [])

        # Auto-infer layout if missing
        if not layout:
            if "stats" in slide_data and len(slide_data["stats"]) > 0:
                layout = "stats"
            elif "comparison" in slide_data:
                layout = "comparison"
            elif "steps" in slide_data and len(slide_data["steps"]) > 0:
                layout = "steps"
            elif "cards" in slide_data and len(slide_data["cards"]) > 0:
                layout = "cards"
            elif "quote" in slide_data:
                layout = "quote"
            else:
                cycle = ["stats", "cards", "comparison", "steps", "quote", "agenda"]
                layout = cycle[(index - 1) % len(cycle)]

        # ============================================================
        # 3. LAYOUT STATS / KPI CARDS (ESTILO CONSULTORÍA EJECUTIVA)
        # ============================================================
        if layout == "stats":
            stats = slide_data.get("stats", [])
            if not stats:
                stats = [
                    {"value": "-35%", "label": "Consumo Energético", "desc": "Reducción directa de demanda en procesos clave."},
                    {"value": "100%", "label": "Privacidad de Datos", "desc": "Procesamiento local confidencial sin salida externa."},
                    {"value": "2.8x", "label": "Agilidad Operativa", "desc": "Aceleración de análisis y toma de decisiones."},
                ]
            for si, stat in enumerate(stats[:3]):
                left_s = 0.9 + si * 3.95
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_s), Inches(1.95), Inches(3.65), Inches(4.5))
                card.fill.solid()
                card.fill.fore_color.rgb = color(theme["card"])
                card.line.color.rgb = color(theme["secondary"])

                # Top Accent Tab
                tab = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left_s), Inches(1.95), Inches(3.65), Inches(0.12))
                tab.fill.solid()
                tab.fill.fore_color.rgb = color(theme["accent"])
                tab.line.fill.background()

                # Stat Label Badge
                s_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_s + 0.3), Inches(2.25), Inches(3.05), Inches(0.3))
                s_badge.fill.solid()
                s_badge.fill.fore_color.rgb = color(theme["secondary"])
                s_badge.line.fill.background()
                add_text(slide, str(stat.get("label", "MÉTRICA CLAVE")).upper(), left_s + 0.3, 2.28, 3.05, 0.24, 9, theme["accent"], True, align=PP_ALIGN.CENTER)

                # Giant Value
                add_text(slide, str(stat.get("value", "99%")), left_s + 0.2, 2.7, 3.25, 1.2, 46, theme["ink"], True, font="Georgia", align=PP_ALIGN.CENTER)

                # Subtitle Highlight
                add_text(slide, "Impacto Cuantitativo", left_s + 0.3, 3.95, 3.05, 0.4, 12.5, theme["highlight"], True, align=PP_ALIGN.CENTER)

                # Description
                add_text(slide, str(stat.get("desc", "")), left_s + 0.35, 4.45, 2.95, 1.6, 11.5, theme["muted"], False, align=PP_ALIGN.CENTER)

        # ============================================================
        # 4. LAYOUT COMPARISON / RETO VS SOLUCIÓN (2 Columnas)
        # ============================================================
        elif layout == "comparison":
            # Left Card: Reto
            card_l = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.95), Inches(5.6), Inches(4.5))
            card_l.fill.solid()
            card_l.fill.fore_color.rgb = color(theme["card"])
            card_l.line.color.rgb = color(theme["secondary"])

            strip_l = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.95), Inches(5.6), Inches(0.12))
            strip_l.fill.solid()
            strip_l.fill.fore_color.rgb = color(theme["highlight"])
            strip_l.line.fill.background()

            add_text(slide, "SITUACIÓN ACTUAL / BARRERAS DETECTADAS", 1.2, 2.25, 5.0, 0.3, 10, theme["highlight"], True)
            add_text(slide, "Retos y Limitaciones", 1.2, 2.6, 5.0, 0.5, 18, theme["ink"], True)

            l_bullets = bullets[:2] if bullets else ["Costes elevados de operación y dependencia externa.", "Falta de integración con datos propios en tiempo real."]
            for bi, b in enumerate(l_bullets):
                top_b = 3.25 + bi * 1.35
                dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), Inches(top_b + 0.05), Inches(0.2), Inches(0.2))
                dot.fill.solid()
                dot.fill.fore_color.rgb = color(theme["highlight"])
                dot.line.fill.background()
                add_text(slide, str(b), 1.55, top_b, 4.7, 1.2, bullet_size, theme["muted"])

            # Right Card: Solución CIRCE
            card_r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.95), Inches(5.6), Inches(4.5))
            card_r.fill.solid()
            card_r.fill.fore_color.rgb = color(theme["card"])
            card_r.line.color.rgb = color(theme["secondary"])

            strip_r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.85), Inches(1.95), Inches(5.6), Inches(0.12))
            strip_r.fill.solid()
            strip_r.fill.fore_color.rgb = color(theme["accent"])
            strip_r.line.fill.background()

            add_text(slide, "PROPUESTA TECNOLÓGICA & VALOR CIRCE", 7.15, 2.25, 5.0, 0.3, 10, theme["accent"], True)
            add_text(slide, "Solución y Resultados", 7.15, 2.6, 5.0, 0.5, 18, theme["ink"], True)

            r_bullets = bullets[2:4] if len(bullets) >= 4 else ["Automatización con modelos locales altamente eficientes.", "Garantía de privacidad, soberanía de datos y bajo consumo."]
            for bi, b in enumerate(r_bullets):
                top_b = 3.25 + bi * 1.35
                dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.15), Inches(top_b + 0.05), Inches(0.2), Inches(0.2))
                dot.fill.solid()
                dot.fill.fore_color.rgb = color(theme["accent"])
                dot.line.fill.background()
                add_text(slide, str(b), 7.5, top_b, 4.7, 1.2, bullet_size, theme["ink"], True)

        # ============================================================
        # 5. LAYOUT STRATEGIC PILLARS / CARDS (3 Columnas con Cabecera)
        # ============================================================
        elif layout == "cards":
            cards = slide_data.get("cards", [])
            if not cards:
                b_slice = bullets[:3] if bullets else ["Optimización y sensorización", "Modelado predictivo", "Implantación segura"]
                cards = [{"title": f"Pilar {ci+1}", "bullets": [b]} for ci, b in enumerate(b_slice)]

            num_cards = min(3, len(cards))
            card_w = 3.65 if num_cards == 3 else 5.6
            spacing = 3.95 if num_cards == 3 else 5.95
            for ci, cdata in enumerate(cards[:num_cards]):
                left_c = 0.9 + ci * spacing
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_c), Inches(1.95), Inches(card_w), Inches(4.5))
                card.fill.solid()
                card.fill.fore_color.rgb = color(theme["card"])
                card.line.color.rgb = color(theme["secondary"])

                # Header Banner
                header_ban = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left_c), Inches(1.95), Inches(card_w), Inches(0.65))
                header_ban.fill.solid()
                header_ban.fill.fore_color.rgb = color(theme["accent"])
                header_ban.line.fill.background()

                add_text(slide, str(cdata.get("title", f"Pilar {ci+1}")).upper(), left_c + 0.25, 2.12, card_w - 0.5, 0.35, 11, "FFFFFF", True)

                # Card Content
                c_bullets = cdata.get("bullets", [])
                for bi, b in enumerate(c_bullets[:3]):
                    top_b = 2.85 + bi * 1.15
                    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left_c + 0.25), Inches(top_b + 0.05), Inches(0.18), Inches(0.18))
                    dot.fill.solid()
                    dot.fill.fore_color.rgb = color(theme["highlight"])
                    dot.line.fill.background()
                    add_text(slide, str(b), left_c + 0.55, top_b, card_w - 0.8, 1.0, bullet_size, theme["ink"], True)

        # ============================================================
        # 6. LAYOUT ROADMAP / FASES (Proceso Secuencial)
        # ============================================================
        elif layout == "steps":
            steps = slide_data.get("steps", [])
            if not steps:
                steps = [
                    {"step": "01", "title": "Diagnóstico Inicial", "desc": "Mapeo exhaustivo de variables, datos y requisitos."},
                    {"step": "02", "title": "Validación & Piloto", "desc": "Implantación en banco de pruebas y ajuste del modelo."},
                    {"step": "03", "title": "Escalado Operativo", "desc": "Despliegue integral con medición continua de KPIs."},
                ]
            num_steps = min(3, len(steps))
            step_w = 3.65
            spacing = 3.95

            # Connection Line
            conn = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.35), Inches(10.8), Inches(0.04))
            conn.fill.solid()
            conn.fill.fore_color.rgb = color(theme["accent"])
            conn.line.fill.background()

            for si, step in enumerate(steps[:num_steps]):
                left_st = 0.9 + si * spacing
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left_st), Inches(1.95), Inches(step_w), Inches(4.5))
                card.fill.solid()
                card.fill.fore_color.rgb = color(theme["card"])
                card.line.color.rgb = color(theme["secondary"])

                # Milestone Oval
                mile = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left_st + 0.3), Inches(2.15), Inches(0.55), Inches(0.55))
                mile.fill.solid()
                mile.fill.fore_color.rgb = color(theme["accent"])
                mile.line.fill.background()
                add_text(slide, f"0{si+1}", left_st + 0.3, 2.2, 0.55, 0.45, 12, "FFFFFF", True, align=PP_ALIGN.CENTER)

                add_text(slide, f"FASE 0{si+1}", left_st + 1.0, 2.25, 2.3, 0.35, 10, theme["highlight"], True)
                add_text(slide, str(step.get("title", f"Fase {si+1}")), left_st + 0.3, 2.95, step_w - 0.6, 0.7, 16, theme["ink"], True)
                add_text(slide, str(step.get("desc", "")), left_st + 0.3, 3.8, step_w - 0.6, 2.2, bullet_size + 1, theme["muted"])

        # ============================================================
        # 7. LAYOUT QUOTE / EXECUTIVE TAKEAWAY (Conclusión Hero)
        # ============================================================
        elif layout == "quote":
            # Dark Hero Container
            panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.95), Inches(11.55), Inches(4.5))
            panel.fill.solid()
            panel.fill.fore_color.rgb = color(theme["ink"])
            panel.line.fill.background()

            add_text(slide, "“", 1.3, 2.05, 1.0, 0.9, 64, theme["accent"], True, font="Georgia")

            quote_text = slide_data.get("quote") or (bullets[0] if bullets else "La innovación tecnológica aplicada es el motor clave para transformar la sostenibilidad en una ventaja competitiva.")
            add_text(slide, str(quote_text), 1.5, 3.0, 10.3, 1.8, 22, "FFFFFF", True, font="Georgia")

            pill_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(5.15), Inches(10.3), Inches(0.8))
            pill_box.fill.solid()
            pill_box.fill.fore_color.rgb = color(theme["accent"])
            pill_box.line.fill.background()

            takeaway_msg = "CONCLUSIÓN EJECUTIVA: " + (bullets[1] if len(bullets) > 1 else "Evidencia cuantitativa respaldada por los modelos de Fundación CIRCE.")
            add_text(slide, takeaway_msg, 1.7, 5.35, 9.9, 0.45, 12, "FFFFFF", True)

        # ============================================================
        # 8. LAYOUT AGENDA / ESTRUCTURA
        # ============================================================
        else:
            items = slide_data.get("items") or bullets or ["Contexto y Diagnóstico", "Metodología y Tecnología", "Resultados Esperados", "Próximos Pasos"]
            for ai, item in enumerate(items[:4]):
                top_a = 2.0 + ai * 1.1
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(top_a), Inches(11.55), Inches(0.9))
                card.fill.solid()
                card.fill.fore_color.rgb = color(theme["card"])
                card.line.color.rgb = color(theme["secondary"])

                pill = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), Inches(top_a + 0.18), Inches(0.54), Inches(0.54))
                pill.fill.solid()
                pill.fill.fore_color.rgb = color(theme["accent"])
                pill.line.fill.background()
                add_text(slide, f"0{ai+1}", 1.2, top_a + 0.22, 0.54, 0.45, 12, "FFFFFF", True, align=PP_ALIGN.CENTER)

                add_text(slide, str(item), 2.05, top_a + 0.22, 10.1, 0.55, 15, theme["ink"], True)

        # Bottom Slide Footer
        bot_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(6.8), Inches(11.55), Inches(0.02))
        bot_line.fill.solid()
        bot_line.fill.fore_color.rgb = color(theme["secondary"])
        bot_line.line.fill.background()

        add_text(slide, "FUNDACIÓN CIRCE · CENTRO TECNOLÓGICO · Conocimiento que se convierte en acción", 0.9, 6.85, 8.5, 0.25, 8, theme["muted"])
        add_text(slide, "Documento Confidencial · Procesado 100% Local", 8.5, 6.85, 3.9, 0.25, 8, theme["muted"], align=PP_ALIGN.RIGHT)

    safe_topic = re.sub(r"[^a-zA-Z0-9áéíóúÁÉÍÓÚñÑ_-]+", "-", topic).strip("-").lower()[:50] or "presentacion"
    filename = f"{safe_topic}-{int(time.time())}.pptx"
    presentation.save(os.path.join(OUTPUT_DIR, filename))
    return filename


def presentation_plan(message: str, preferences: dict, document_context: dict | None = None) -> list[dict]:
    num_slides_pref = int(preferences.get("num_slides", 7))

    fallback = [
        {"layout": "cover", "title": message[:70], "subtitle": "Estrategia ejecutiva y análisis cuantitativo para CIRCE"},
        {"layout": "agenda", "title": "Agenda y Enfoque Estratégico", "items": ["Diagnóstico y Contexto", "Solución Tecnológica", "Métricas y Retorno Cuantitativo", "Plan de Acción y Cronograma"]},
        {
            "layout": "stats",
            "title": "Métricas y Rendimiento Esperado",
            "subtitle": "Indicadores cuantitativos de impacto directo en la operación",
            "stats": [
                {"value": "-35%", "label": "Demanda Energética", "desc": "Reducción directa estimada mediante optimización local."},
                {"value": "100%", "label": "Soberanía de Datos", "desc": "Privacidad y confidencialidad total sin exposición externa."},
                {"value": "2.8x", "label": "Productividad", "desc": "Agilización en análisis documental y generación de informes."},
            ],
        },
        {
            "layout": "comparison",
            "title": "Comparativa de Escenarios",
            "subtitle": "Evaluación del estado actual frente a la propuesta tecnológica CIRCE",
            "bullets": [
                "Procesos manuales y dependencia de servicios en la nube con coste recurrente.",
                "Riesgo de fuga de información corporativa y latencia en consultas técnicas.",
                "Arquitectura local de alta concurrencia con semáforos de capacidad.",
                "Generación automática de entregables ejecutivos en Word y PowerPoint.",
            ],
        },
        {
            "layout": "cards",
            "title": "Ejes Estratégicos de Innovación",
            "subtitle": "Pilares fundamentales para la adopción tecnológica en la organización",
            "cards": [
                {"title": "Arquitectura Local", "bullets": ["Inferencia GGUF cuantizada", "Concurrencia multihilo"]},
                {"title": "Gobernanza & Seguridad", "bullets": ["Sin telemetría externa", "Conexión segura vía túnel"]},
                {"title": "Impacto Operativo", "bullets": ["Documentos Word corporativos", "Diapositivas PPTX ejecutivas"]},
            ],
        },
        {
            "layout": "steps",
            "title": "Plan de Implementación y Fases",
            "subtitle": "Cronograma de ejecución estructurado para mitigar riesgos",
            "steps": [
                {"step": "01", "title": "Diagnóstico y Requisitos", "desc": "Mapeo de datos y validación de casos de uso prioritarios."},
                {"step": "02", "title": "Validación en Piloto", "desc": "Despliegue en banco de pruebas con equipos técnicos clave."},
                {"step": "03", "title": "Escalado Transversal", "desc": "Adopción operativa con monitorización y métricas en tiempo real."},
            ],
        },
        {"kind": "closing", "layout": "closing", "title": "Conclusiones y Próximos Pasos", "subtitle": "Preguntas, debate y calendario de hitos"},
    ]

    if DEMO_MODE:
        return fallback[:num_slides_pref] if num_slides_pref >= 4 else fallback

    doc_extra = ""
    if document_context and document_context.get("text"):
        doc_extra = (
            f"Basate estrictamente en el siguiente documento adjunto ('{document_context.get('filename')}'):\n"
            f"{document_context['text'][:12000]}\n\n"
        )

    prompt = (
        "Devuelve SOLO JSON valido, sin markdown ni texto extra, con este esquema exacto:\n"
        '{"slides":[\n'
        '  {"layout":"cover","title":"...","subtitle":"..."},\n'
        '  {"layout":"agenda","title":"Agenda","items":["...","...","..."]},\n'
        '  {"layout":"stats","title":"Metricas Clave","subtitle":"...","stats":[{"value":"...","label":"...","desc":"..."}]},\n'
        '  {"layout":"comparison","title":"Reto vs Solucion","subtitle":"...","bullets":["...","...","...","..."]},\n'
        '  {"layout":"cards","title":"Ejes Tecnologicos","cards":[{"title":"...","bullets":["...","..."]}]},\n'
        '  {"layout":"steps","title":"Hoja de Ruta","steps":[{"step":"01","title":"...","desc":"..."}]},\n'
        '  {"kind":"closing","layout":"closing","title":"Cierre y Proximos Pasos","subtitle":"..."}\n'
        ']}\n'
        f"Genera exactamente {num_slides_pref} diapositivas de nivel ejecutivo para consultoría estratégica sobre este tema: '{message}'.\n"
        f"Tono: {preferences.get('tone', 'profesional')}. Densidad: {preferences.get('density', 'medio')}.\n"
        f"{doc_extra}"
    )

    try:
        raw = ollama_chat(prompt, "Generación avanzada de presentaciones", PRESENTATION_MODEL)
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end != -1:
            parsed = json.loads(raw[start : end + 1])
            slides = parsed.get("slides", [])
            if slides and all(isinstance(s, dict) for s in slides):
                if slides[-1].get("kind") != "closing" and slides[-1].get("layout") != "closing":
                    slides.append({"kind": "closing", "layout": "closing", "title": "Cierre", "subtitle": "Preguntas y próximos pasos"})
                return slides
    except Exception:
        pass
    return fallback


def generate_presentation(message: str, preferences: dict, document_context: dict | None = None) -> tuple[str, str]:
    slides = presentation_plan(message, preferences, document_context)
    filename = create_presentation(message, slides, preferences)
    doc_note = f" a partir del documento '{document_context['filename']}'" if document_context else ""
    return (
        filename,
        f"He diseñado y generado una presentación PowerPoint ejecutiva (.pptx) con {len(slides)} diapositivas{doc_note}. "
        f"Aplica diseño de consultoría estratégica, paleta {preferences.get('palette', 'circe')}, infografías, métricas clave y tipografía institucional.",
    )


# ==========================================
# SERVIDOR HTTP Y CONTROLADOR REST
# ==========================================

class Handler(BaseHTTPRequestHandler):
    def send_cors_headers(self) -> None:
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type, X-Client-Id")

    def send_json(self, status: int, data: dict) -> None:
        body = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_cors_headers()
        self.end_headers()
        self.wfile.write(body)

    def do_OPTIONS(self) -> None:
        self.send_response(204)
        self.send_cors_headers()
        self.end_headers()

    def do_GET(self) -> None:
        parsed_url = urlparse(self.path)
        path = parsed_url.path
        query = parse_qs(parsed_url.query)

        client_id = self.headers.get("X-Client-Id") or (query.get("client_id", [None])[0])
        if client_id:
            metrics_tracker.touch_user(client_id)

        if path == "/api/health":
            summary = metrics_tracker.get_summary()["summary"]
            self.send_json(
                200,
                {
                    "status": "ok",
                    "model": MODEL_NAME,
                    "mode": "demo" if DEMO_MODE else "ollama",
                    "active_users": summary["active_users_5m"],
                    "active_requests": summary["active_requests"],
                    "max_concurrency_slots": MAX_CONCURRENT_REQUESTS,
                    "ollama_parallel_slots": OLLAMA_NUM_PARALLEL,
                },
            )
            return

        if path == "/api/metrics":
            data = metrics_tracker.get_summary()
            self.send_json(200, data)
            return

        if path.startswith("/generated/"):
            filename = os.path.basename(path.removeprefix("/generated/"))
            file_path = os.path.join(OUTPUT_DIR, filename)
            if os.path.isfile(file_path):
                with open(file_path, "rb") as generated_file:
                    body = generated_file.read()

                ext = os.path.splitext(filename.lower())[1]
                if ext == ".pptx":
                    content_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                elif ext == ".docx":
                    content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                elif ext == ".pdf":
                    content_type = "application/pdf"
                else:
                    content_type = "application/octet-stream"

                self.send_response(200)
                self.send_header("Content-Type", content_type)
                self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
                self.send_header("Content-Length", str(len(body)))
                self.send_cors_headers()
                self.end_headers()
                self.wfile.write(body)
                return
            self.send_json(404, {"error": "file not found"})
            return

        # Serve static assets from Vite dist/ build if available
        if os.path.isdir(DIST_DIR):
            relative_path = path.lstrip("/")
            file_path = os.path.join(DIST_DIR, relative_path)
            if not os.path.isfile(file_path):
                file_path = os.path.join(DIST_DIR, "index.html")

            if os.path.isfile(file_path):
                mime_type, _ = mimetypes.guess_type(file_path)
                if not mime_type:
                    mime_type = "application/octet-stream"
                try:
                    with open(file_path, "rb") as static_file:
                        content = static_file.read()
                    self.send_response(200)
                    self.send_header("Content-Type", mime_type)
                    self.send_header("Content-Length", str(len(content)))
                    self.send_cors_headers()
                    self.end_headers()
                    self.wfile.write(content)
                    return
                except Exception as e:
                    self.send_json(500, {"error": "Error reading static file", "detail": str(e)})
                    return

        self.send_json(404, {"error": "not found"})

    def do_POST(self) -> None:
        if self.path == "/api/heartbeat":
            size = int(self.headers.get("Content-Length", "0"))
            client_id = None
            if size > 0:
                try:
                    body = json.loads(self.rfile.read(size))
                    client_id = body.get("client_id")
                except Exception:
                    pass
            client_id = client_id or self.headers.get("X-Client-Id")
            if client_id:
                metrics_tracker.touch_user(client_id)
            self.send_json(200, {"status": "ok", "client_id": client_id})
            return

        # ----------------------------------------------------
        # SUBIDA Y EXTRACCIÓN DE DOCUMENTOS (PDF, DOCX, TXT)
        # ----------------------------------------------------
        if self.path == "/api/upload-document":
            size = int(self.headers.get("Content-Length", "0"))
            if size == 0:
                self.send_json(400, {"error": "No file content received"})
                return

            raw_body = self.rfile.read(size)
            try:
                data = json.loads(raw_body)
                filename = data.get("filename", "documento.txt")
                mime_type = data.get("mime_type", "")
                base64_content = data.get("content_base64", "")
                if base64_content:
                    if "," in base64_content:
                        base64_content = base64_content.split(",", 1)[1]
                    file_bytes = base64.b64decode(base64_content)
                else:
                    file_bytes = b""
            except Exception:
                filename = "documento.txt"
                mime_type = "text/plain"
                file_bytes = raw_body

            doc_info = extract_document_text(file_bytes, filename, mime_type)
            with metrics_tracker.lock:
                metrics_tracker.total_documents += 1
                metrics_tracker.use_case_counts["Documentos e Informes"] += 1

            self.send_json(200, doc_info)
            return

        # ----------------------------------------------------
        # EXPORTAR DIRECTO A DOCUMENTO WORD (.DOCX)
        # ----------------------------------------------------
        if self.path == "/api/export-docx":
            size = int(self.headers.get("Content-Length", "0"))
            try:
                data = json.loads(self.rfile.read(size))
                title = str(data.get("title", "Informe Técnico CIRCE"))
                content = str(data.get("content", ""))
                topic = str(data.get("topic", "Investigación y Desarrollo"))
                metadata = data.get("metadata", {})
                filename = create_docx_report(title, content, topic, metadata)
                self.send_json(200, {"filename": filename, "download_url": f"/generated/{filename}"})
                return
            except Exception as e:
                self.send_json(500, {"error": "Error generando documento Word", "detail": str(e)})
                return

        # ----------------------------------------------------
        # CHAT PRINCIPAL CON CAPACIDAD MULTIUSUARIO Y MULTI-TURNO
        # ----------------------------------------------------
        if self.path != "/api/chat":
            self.send_json(404, {"error": "not found"})
            return

        # Check concurrency capacity
        acquired = concurrency_semaphore.acquire(timeout=25)
        if not acquired:
            self.send_json(
                429,
                {
                    "error": "Capacidad máxima de concurrencia alcanzada",
                    "detail": f"Hay {MAX_CONCURRENT_REQUESTS} peticiones procesándose en paralelo. Por favor, reintenta en unos instantes.",
                },
            )
            return

        client_id = self.headers.get("X-Client-Id", "anon")
        started = time.perf_counter()
        req_id = 0
        use_case = "Asistente general"
        is_presentation = False
        mode = "demo" if DEMO_MODE else "ollama"

        try:
            size = int(self.headers.get("Content-Length", "0"))
            data = json.loads(self.rfile.read(size))
            client_id = str(data.get("client_id") or client_id)
            req_id, _ = metrics_tracker.start_request(client_id)

            message = str(data.get("message", "")).strip()
            messages_history = data.get("messages", [])
            presentation_topic = str(data.get("presentation_topic", "")).strip()
            presentation_preferences = data.get("presentation_preferences") or {}
            document_context = data.get("document_context")
            use_case = str(data.get("use_case", "Asistente general"))

            if not message:
                latency_ms = round((time.perf_counter() - started) * 1000)
                metrics_tracker.end_request(req_id, client_id, use_case, latency_ms, 400, mode, False)
                self.send_json(400, {"error": "message is required"})
                return

            # Extract last assistant content from conversation history if user refers to "esto"
            last_assistant_content = ""
            last_user_query = ""
            for prev in reversed(messages_history[:-1] if messages_history and messages_history[-1].get("role") == "user" else messages_history):
                if prev.get("role") == "assistant" and prev.get("content") and not last_assistant_content:
                    last_assistant_content = prev.get("content", "")
                elif prev.get("role") == "user" and prev.get("content") and not last_user_query:
                    last_user_query = prev.get("content", "")
                if last_assistant_content and last_user_query:
                    break

            # Automatic preference extraction from prompt if user gave hints
            pref_palette = presentation_preferences.get("palette")
            if not pref_palette:
                m_low = message.lower()
                if "azul" in m_low:
                    pref_palette = "azul"
                elif "cálida" in m_low or "calida" in m_low or "naranja" in m_low:
                    pref_palette = "calida"
                elif "púrpura" in m_low or "purpura" in m_low or "violeta" in m_low:
                    pref_palette = "purpura"
                elif "monocroma" in m_low or "gris" in m_low or "grafito" in m_low:
                    pref_palette = "monocroma"
                else:
                    pref_palette = "circe"
                presentation_preferences["palette"] = pref_palette

            if "num_slides" not in presentation_preferences:
                m_low = message.lower()
                if "5" in m_low or "breve" in m_low or "corta" in m_low:
                    presentation_preferences["num_slides"] = 5
                elif "10" in m_low or "detallada" in m_low or "larga" in m_low:
                    presentation_preferences["num_slides"] = 10
                else:
                    presentation_preferences["num_slides"] = 7

            if "tone" not in presentation_preferences:
                presentation_preferences["tone"] = "profesional"
            if "density" not in presentation_preferences:
                presentation_preferences["density"] = "medio"

            download_url = None
            docx_download_url = None

            # ----------------------------------------------------
            # CASO 1: CONVERSATIONAL PRESENTATION REQUEST
            # ----------------------------------------------------
            if presentation_topic or is_presentation_intent(message):
                is_presentation = True
                # Topic resolution: if user says "conviértelo a presentación", use previous assistant response or document
                topic_source = presentation_topic or message
                if len(message.split()) <= 6 and ("esto" in message.lower() or "anterior" in message.lower() or "documento" in message.lower()):
                    if document_context:
                        topic_source = f"Análisis de {document_context.get('filename')}"
                    elif last_user_query:
                        topic_source = last_user_query
                    elif last_assistant_content:
                        first_line = last_assistant_content.splitlines()[0].replace("#", "").strip()
                        topic_source = first_line[:80] or "Presentación Ejecutiva CIRCE"

                filename, reply = generate_presentation(topic_source, presentation_preferences, document_context)
                download_url = f"/generated/{filename}"

            # ----------------------------------------------------
            # CASO 2: CONVERSATIONAL WORD DOCX EXPORT REQUEST
            # ----------------------------------------------------
            elif is_export_word_intent(message):
                # If user says "expórtalo a word" or "pásalo a docx" referring to the last message
                is_short_command = len(message.split()) <= 6 and any(w in message.lower() for w in ("esto", "anterior", "informe", "word", "docx"))

                if is_short_command and last_assistant_content:
                    doc_title = "Informe Técnico CIRCE"
                    first_line = last_assistant_content.splitlines()[0].replace("#", "").strip()
                    if first_line and len(first_line) > 3:
                        doc_title = first_line[:60]
                    docx_fn = create_docx_report(
                        title=doc_title,
                        markdown_content=last_assistant_content,
                        topic=use_case,
                        metadata={"doc_name": document_context.get("filename") if document_context else None},
                    )
                    docx_download_url = f"/generated/{docx_fn}"
                    reply = (
                        "He preparado el **documento formal en formato Word (.docx)** con el contenido del informe solicitado.\n\n"
                        "Aplica el formato institucional de Fundación CIRCE con caja de metadatos, tablas y tipografía corporativa. Puedes descargarlo directamente a continuación:"
                    )
                else:
                    # User asked to generate a new report and export it
                    if DEMO_MODE:
                        doc_name_str = f" a partir de '{document_context['filename']}'" if document_context else ""
                        reply = (
                            f"# Informe Técnico: {message[:50]}\n\n"
                            f"**Área:** {use_case} · **Fundación CIRCE**\n\n"
                            f"Este informe ha sido elaborado{doc_name_str} siguiendo la metodología de análisis local.\n\n"
                            "## 1. Resumen Ejecutivo\n\n"
                            "- Optimización directa de recursos y eficiencia en procesos.\n"
                            "- Privacidad absoluta de los datos sin salida a la nube pública.\n"
                            "- Validación empírica alineada con la estrategia de innovación.\n\n"
                            "## 2. Métricas y Conclusiones\n\n"
                            "- **Retorno estimado:** Alta viabilidad en fase piloto.\n"
                            "- **Recomendación prioritaria:** Despliegue secuencial por áreas técnicas."
                        )
                    else:
                        reply = ollama_chat(message, use_case, MODEL_NAME, document_context, messages_history)

                    docx_fn = create_docx_report(
                        title=f"Informe CIRCE: {message[:45]}",
                        markdown_content=reply,
                        topic=use_case,
                        metadata={"doc_name": document_context.get("filename") if document_context else None},
                    )
                    docx_download_url = f"/generated/{docx_fn}"

            # ----------------------------------------------------
            # CASO 3: CONSULTA GENERAL O ANÁLISIS DE DOCUMENTO
            # ----------------------------------------------------
            elif DEMO_MODE:
                doc_note = f" He procesado el documento adjunto '{document_context['filename']}' ({document_context.get('word_count', 0)} palabras)." if document_context else ""
                reply = (
                    f"He recibido tu consulta sobre {use_case.lower()}.{doc_note}\n\n"
                    "*(Modo demo activo: conecta Ollama y cambia `DEMO_MODE=false` en `.env` para inferencia local con Hugging Face)*.\n\n"
                    "Puedes pedirme en cualquier momento:\n"
                    "- *\"Expórtalo a Word\"* para obtener un informe formateado descargable.\n"
                    "- *\"Conviértelo en presentación\"* para generar diapositivas PPTX ejecutivas."
                )
            else:
                reply = ollama_chat(message, use_case, MODEL_NAME, document_context, messages_history)

            latency_ms = round((time.perf_counter() - started) * 1000)
            metrics_tracker.end_request(req_id, client_id, use_case, latency_ms, 200, mode, is_presentation)

            response = {
                "reply": reply,
                "model": MODEL_NAME,
                "mode": mode,
                "latency_ms": latency_ms,
                "client_id": client_id,
            }
            if download_url:
                response["download_url"] = download_url
            if docx_download_url:
                response["docx_download_url"] = docx_download_url

            self.send_json(200, response)

        except Exception as error:
            latency_ms = round((time.perf_counter() - started) * 1000)
            if req_id > 0:
                metrics_tracker.end_request(req_id, client_id, use_case, latency_ms, 502, mode, is_presentation)
            self.send_json(502, {"error": "model unavailable", "detail": str(error)})

        finally:
            concurrency_semaphore.release()


if __name__ == "__main__":
    print("=" * 65)
    print(" CIRCE Local AI Lab - Servidor Concurrente Inteligente")
    print(f" URL local:             http://{HOST}:{PORT}")
    print(f" Modo:                  {'DEMO' if DEMO_MODE else 'OLLAMA LOCAL'}")
    print(f" Modelo Chat / Docs:    {MODEL_NAME}")
    print(f" Modelo Presentaciones: {PRESENTATION_MODEL}")
    print(f" Concurrencia Max:      {MAX_CONCURRENT_REQUESTS} slots paralelos")
    print(f" Ollama Parallel:       {OLLAMA_NUM_PARALLEL} workers")
    print(f" Procesado Documentos:  PDF, DOCX, TXT, MD soportados")
    print(f" Generación PPTX/DOCX:  Plantillas Ejecutivas CIRCE integradas")
    if os.path.isdir(DIST_DIR):
        print(f" Frontend:              Servido automáticamente desde dist/")
    print("=" * 65)
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
